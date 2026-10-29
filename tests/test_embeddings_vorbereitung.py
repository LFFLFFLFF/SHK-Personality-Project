"""Gezielte Tests der Datenprüfung ohne Modell, Download oder echte Embeddings."""

from contextlib import redirect_stderr, redirect_stdout
import csv
import importlib.util
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("stage1_embeddings", ROOT / "01_embeddings.py")
stage1 = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(stage1)


def add_table(sheet, name, ref):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


class WorkbookValidationTests(unittest.TestCase):
    def setUp(self):
        # Testdaten bleiben im freigegebenen Projektordner und werden danach entfernt.
        self.temp = tempfile.TemporaryDirectory(prefix="test_tmp_", dir=ROOT)
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.input = self.root / "data" / "input.xlsx"
        self.input.parent.mkdir()
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Items"
        sheet.append(stage1.ITEM_HEADERS)
        sheet.append(["A1", "Ich denke häufig über Vergangenes nach.", None,
            "Fragebogen A", "Rumination", "direkt", "Vergangene Dinge beschäftigen mich häufig.",
            "Ich grüble oft über vergangene Dinge.", "Mit Vergangenem befasse ich mich gedanklich oft.",
            "Anmerkung", "direkt"])
        sheet.append(["B1inv.", "Ich vertraue Menschen grundsätzlich.", "Ich misstraue Menschen grundsätzlich.",
            "Fragebogen B", "Hostility", "invertiert", "Anderen Menschen begegne ich mit Misstrauen.",
            "Grundsätzlich bin ich Menschen gegenüber misstrauisch.", "Menschen misstraue ich im Allgemeinen.",
            "Textuell invertiert", "direkt"])
        add_table(sheet, "ItemsAugmentierung", "A1:K3")
        controls = workbook.create_sheet("Kontrolltexte")
        controls.append(stage1.CONTROL_HEADERS)
        controls.append(["KONT01", "kontrolle", "Alle meine Entchen schwimmen auf dem See.", "Separater Kontrolltext"])
        add_table(controls, "KontrolltexteTabelle", "A1:D2")
        workbook.save(self.input)
        workbook.close()
        self.config = {
            "schema_version": 1, "input_file": "data/input.xlsx", "item_sheet": "Items",
            "item_table": "ItemsAugmentierung", "control_sheet": "Kontrolltexte",
            "control_table": "KontrolltexteTabelle", "expected_items": 2,
            "expected_variants": 6, "expected_regular_texts": 8,
            "expected_construct_counts": {"Rumination": 1, "Hostility": 1},
            "expected_source_counts": {"B": 1, "C": 1}, "expected_control_count": 1,
            "forbidden_item_ids": ["RRQ3"], "expected_inverted_item_ids": ["B1inv."],
            "model_name": "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2",
            "model_revision": "e8f8c211226b894fcb81acc59f3b34ba3efd5f42",
            "expected_embedding_dimension": 384, "embedding_prompt": "", "batch_size": 16,
            "device": "cpu", "output_dir": "runs", "provenance": {"source_file": "Itemselection.py"},
        }
        self.config_path = self.root / "config_embeddings.json"
        self.save_config()

    def save_config(self):
        self.config_path.write_text(json.dumps(self.config, ensure_ascii=False), encoding="utf-8")

    def mutate(self, callback):
        workbook = openpyxl.load_workbook(self.input)
        callback(workbook)
        workbook.save(self.input)
        workbook.close()

    def validate(self):
        return stage1.validate_input(self.config, self.root)

    def assert_invalid(self, expected_message):
        report, records = self.validate()
        self.assertFalse(report["data_suitable"])
        self.assertIn(expected_message, " ".join(report["errors"]))
        self.assertFalse(report["model_loaded"])
        self.assertFalse(report["embeddings_computed"])
        return report, records

    def test_validation_succeeds_without_importing_model_or_numpy(self):
        original_import = __import__

        def guarded_import(name, *args, **kwargs):
            if name.split(".")[0] in {"torch", "sentence_transformers", "transformers", "numpy"}:
                self.fail(f"Unzulässiger Import beim Prüflauf: {name}")
            return original_import(name, *args, **kwargs)

        with patch("builtins.__import__", side_effect=guarded_import), \
             patch.object(stage1, "compute_embeddings", side_effect=AssertionError("Kein Modell im Prüflauf")):
            report, records = self.validate()
        self.assertTrue(report["data_suitable"], report["errors"])
        self.assertEqual(report["counts"]["regular_texts"], 8)
        self.assertEqual(report["counts"]["control_texts"], 1)
        self.assertEqual(report["counts"]["source_counts"], {"B": 1, "C": 1})
        self.assertEqual(len(report["source_fingerprints"]), 2)
        self.assertEqual(len(report["source_sha256"]), 64)
        self.assertFalse(report["model_downloads_allowed"])
        self.assertEqual(len(records), 9)
        self.assertFalse(any(report["stage_flags"].values()))

    def test_blank_working_text_falls_back_only_to_same_rows_original(self):
        self.mutate(lambda w: setattr(w["Items"]["C2"], "value", "  \n  "))
        report, records = self.validate()
        self.assertTrue(report["data_suitable"], report["errors"])
        self.assertEqual(records[0]["source_cell"], "B2")
        self.assertEqual(records[0]["text"], "Ich denke häufig über Vergangenes nach.")

    def test_nonempty_working_text_is_selected_and_inner_text_is_preserved(self):
        self.mutate(lambda w: setattr(w["Items"]["C2"], "value", "  Mein  Text\nbleibt so.  "))
        self.config["expected_source_counts"] = {"B": 0, "C": 2}
        report, records = self.validate()
        self.assertTrue(report["data_suitable"], report["errors"])
        self.assertEqual(records[0]["source_cell"], "C2")
        self.assertEqual(records[0]["text"], "Mein  Text\nbleibt so.")

    def test_inverse_original_is_never_selected_for_any_formulation(self):
        report, records = self.validate()
        self.assertTrue(report["data_suitable"])
        inverse = [r for r in records if r["item_id"] == "B1inv."]
        self.assertEqual(inverse[0]["source_cell"], "C3")
        self.assertTrue(all(r["original_polarity"] == "invertiert" and r["analysis_polarity"] == "direkt" for r in inverse))
        self.assertNotIn("Ich vertraue Menschen grundsätzlich.", [r["text"] for r in records])
        self.mutate(lambda w: setattr(w["Items"]["G3"], "value", "Ich vertraue Menschen grundsätzlich."))
        self.assert_invalid("Gegenpoliger Originaltext")

    def test_inverse_requires_explicit_working_text(self):
        self.mutate(lambda w: setattr(w["Items"]["C3"], "value", None))
        self.assert_invalid("textuell invertierte Arbeitsfassung")

    def test_inverse_set_and_analysis_polarity_are_enforced(self):
        self.config["expected_inverted_item_ids"] = ["B1inv.", "FEHLTinv."]
        self.assert_invalid("Invertierte Item-IDs")
        self.config["expected_inverted_item_ids"] = ["B1inv."]
        self.mutate(lambda w: setattr(w["Items"]["K3"], "value", "invertiert"))
        self.assert_invalid("Polung Analyse muss direkt")

    def test_rts_d8_long_basis_allows_exactly_two_variants(self):
        long_text = " ".join(f"Wort{n}" for n in range(28))

        def change(workbook):
            sheet = workbook["Items"]
            sheet["A2"] = "RTS-D8"
            sheet["B2"] = long_text
            sheet["I2"] = None

        self.mutate(change)
        self.config.update(expected_variants=5, expected_regular_texts=7)
        report, records = self.validate()
        self.assertTrue(report["data_suitable"], report["errors"])
        self.assertEqual(len([r for r in records if r["item_id"] == "RTS-D8"]), 3)
        self.mutate(lambda w: setattr(w["Items"]["I2"], "value", "Eine zusätzliche dritte Variante."))
        self.assert_invalid("genau 2 Varianten")

    def test_short_basis_requires_third_variant(self):
        self.mutate(lambda w: setattr(w["Items"]["I2"], "value", None))
        self.assert_invalid("genau 3 Varianten")

    def test_footer_outside_native_table_is_not_an_item(self):
        self.mutate(lambda w: setattr(w["Items"]["A54"], "value", "54. Fußnote ist kein Kontrollitem"))
        report, records = self.validate()
        self.assertTrue(report["data_suitable"], report["errors"])
        self.assertEqual(report["counts"]["items"], 2)
        self.assertEqual(records[-1]["role"], "kontrolle")
        self.assertEqual(records[-2]["role"], "item")

    def test_control_role_is_explicit_and_can_be_grenzfall(self):
        self.mutate(lambda w: setattr(w["Kontrolltexte"]["B2"], "value", "grenzfall"))
        report, records = self.validate()
        self.assertTrue(report["data_suitable"], report["errors"])
        self.assertEqual(report["counts"]["control_role_counts"], {"grenzfall": 1})
        self.assertEqual(records[-1]["analysis_polarity"], "")
        self.mutate(lambda w: setattr(w["Kontrolltexte"]["B2"], "value", "negativ"))
        self.assert_invalid("Rolle muss explizit")

    def test_last_regular_group_remains_regular_without_controls(self):
        def change(workbook):
            sheet = workbook["Kontrolltexte"]
            sheet.delete_rows(2)
            sheet.tables["KontrolltexteTabelle"].ref = "A1:D1"

        self.mutate(change)
        self.config["expected_control_count"] = 0
        report, records = self.validate()
        self.assertTrue(report["data_suitable"], report["errors"])
        self.assertEqual(report["counts"]["control_texts"], 0)
        self.assertTrue(all(r["role"] == "item" for r in records))

    def test_missing_control_sheet_fails_even_when_last_group_exists(self):
        self.mutate(lambda w: w.remove(w["Kontrolltexte"]))
        self.assert_invalid("Pflichtblatt fehlt")

    def test_missing_input_does_not_create_example_data(self):
        self.config["input_file"] = "fehlend.xlsx"
        report, records = self.assert_invalid("keine Ersatzdaten")
        self.assertEqual(records, [])
        self.assertEqual(report["counts"], {})

    def test_count_mismatch_is_fatal(self):
        self.config["expected_items"] = 53
        self.assert_invalid("Anzahlprüfung items")

    def test_duplicate_id_across_roles_is_fatal(self):
        self.mutate(lambda w: setattr(w["Kontrolltexte"]["A2"], "value", "A1"))
        self.assert_invalid("Doppelte ID zwischen Item-/Kontrollrollen")

    def test_duplicate_item_id_and_forbidden_id_are_fatal(self):
        self.mutate(lambda w: setattr(w["Items"]["A3"], "value", "A1"))
        self.assert_invalid("Doppelte ID")
        self.mutate(lambda w: setattr(w["Items"]["A2"], "value", "RRQ3"))
        self.assert_invalid("Ausgeschlossenes Item")

    def test_duplicate_group_text_is_rejected_without_modifying_text(self):
        self.mutate(lambda w: setattr(w["Items"]["G2"], "value", "  ICH  denke häufig über Vergangenes nach.  "))
        self.assert_invalid("Doppelte Formulierung innerhalb")

    def test_cross_id_duplicate_is_warning_and_is_not_deduplicated(self):
        self.mutate(lambda w: setattr(w["Kontrolltexte"]["C2"], "value", "Ich denke häufig über Vergangenes nach."))
        report, records = self.validate()
        self.assertTrue(report["data_suitable"], report["errors"])
        self.assertEqual(len(records), 9)
        self.assertTrue(any("Keine Deduplizierung" in warning for warning in report["warnings"]))

    def test_formula_error_numeric_and_weird_text_are_rejected(self):
        for address, value, message in (("J2", "=1+1", "Formeln und Excel-Fehler"),
                                        ("B2", "#VALUE!", "Formeln und Excel-Fehler"),
                                        ("B2", 123, "Erwartet wird Text"),
                                        ("B2", "Ungültiges \ufffd Zeichen", "Ersatzzeichen")):
            with self.subTest(address=address, value=value):
                workbook = openpyxl.load_workbook(self.input)
                old = workbook["Items"][address].value
                workbook.close()
                self.mutate(lambda w: setattr(w["Items"][address], "value", value))
                self.assert_invalid(message)
                self.mutate(lambda w: setattr(w["Items"][address], "value", old))

    def test_identifier_and_type_whitespace_are_rejected(self):
        self.mutate(lambda w: setattr(w["Items"]["A2"], "value", " A1"))
        self.assert_invalid("Rand-Leerraum")
        self.mutate(lambda w: setattr(w["Items"]["A2"], "value", "A  1"))
        self.assert_invalid("IDs dürfen nur einfache innere Leerzeichen")
        self.mutate(lambda w: setattr(w["Items"]["A2"], "value", "A1"))
        self.mutate(lambda w: setattr(w["Kontrolltexte"]["B2"], "value", " kontrolle"))
        self.assert_invalid("Rand-Leerraum")

    def test_original_ids_with_one_inner_space_are_preserved(self):
        self.mutate(lambda w: setattr(w["Items"]["A2"], "value", "RC3 36"))
        report, records = self.validate()
        self.assertTrue(report["data_suitable"], report["errors"])
        self.assertEqual(records[0]["item_id"], "RC3 36")

    def test_missing_ml_libraries_are_warnings_not_data_errors(self):
        status = {"sentence-transformers": {"available": False, "version": None}}
        with patch.object(stage1, "library_status", return_value=status):
            report, _ = self.validate()
        self.assertTrue(report["data_suitable"], report["errors"])
        self.assertTrue(any("sentence-transformers" in warning for warning in report["warnings"]))
        self.config["expected_items"] = 99
        with patch.object(stage1, "library_status", return_value=status):
            self.assert_invalid("Anzahlprüfung items")

    def test_mapping_has_stable_separate_array_indices_and_sources(self):
        report, records = self.validate()
        self.assertTrue(report["data_suitable"])
        self.assertEqual([r["array_index"] for r in records[:-1]], list(range(8)))
        self.assertEqual(records[-1]["array_index"], 0)
        self.assertEqual([r["mapping_index"] for r in records], list(range(9)))
        self.assertEqual(records[4]["item_index"], 1)
        self.assertEqual(records[4]["variant_id"], "B1inv.::basis")
        self.assertEqual(records[4]["source_row"], 3)
        self.assertEqual(records[4]["original_source_cell"], "B3")

    def test_default_cli_only_writes_report_and_mapping_when_requested(self):
        with patch.object(stage1, "compute_embeddings", side_effect=AssertionError("Keine Berechnung")), \
             redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            result = stage1.main(["--config", str(self.config_path), "--bericht-dir", "pruefung"])
        self.assertEqual(result, 0)
        directory = self.root / "pruefung"
        self.assertEqual({p.name for p in directory.iterdir()}, {"pruefbericht.json", "text_mapping.csv"})
        report = json.loads((directory / "pruefbericht.json").read_text(encoding="utf-8"))
        self.assertFalse(report["model_loaded"])
        self.assertFalse(report["embeddings_computed"])
        self.assertEqual(len(report["config_sha256"]), 64)
        data = (directory / "text_mapping.csv").read_bytes()
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        with (directory / "text_mapping.csv").open(encoding="utf-8-sig", newline="") as stream:
            self.assertEqual(len(list(csv.DictReader(stream))), 9)
        self.assertFalse((self.root / "runs").exists())

    def test_relative_config_path_is_relative_to_script_not_cwd(self):
        with patch.object(stage1, "SCRIPT_DIR", self.root):
            config, path = stage1.load_config("config_embeddings.json")
        self.assertEqual(path, self.config_path.resolve())
        self.assertEqual(config, self.config)

    def test_invalid_config_revision_and_prompt_stop_before_model_loading(self):
        for key, value, message in (("model_revision", "main", "40 Hexzeichen"),
                                    ("embedding_prompt", "query: ", "explizite leere")):
            with self.subTest(key=key):
                original = self.config[key]
                self.config[key] = value
                self.save_config()
                with self.assertRaisesRegex(stage1.ValidationError, message):
                    stage1.load_config(str(self.config_path))
                self.config[key] = original


class TokenGuardTests(unittest.TestCase):
    def test_oversized_text_is_rejected_before_encode(self):
        class Tokenizer:
            model_max_length = 512

            def __call__(self, texts, **kwargs):
                self.options = kwargs
                return {"input_ids": [list(range(129)) for _ in texts]}

        class FakeModel:
            max_seq_length = 128
            tokenizer = Tokenizer()

            def encode(self, *args, **kwargs):
                raise AssertionError("encode darf bei Überlänge nicht erreicht werden")

        model = FakeModel()
        records = [{"text": "Ein sehr langer Text", "variant_id": "A1::basis"}]
        with self.assertRaisesRegex(stage1.ValidationError, "Keine stille Kürzung"):
            stage1.guard_token_lengths(model, records)
        self.assertFalse(model.tokenizer.options["truncation"])
        self.assertTrue(model.tokenizer.options["add_special_tokens"])
        self.assertEqual(records[0]["token_count"], 129)

    def test_tokenizer_limit_is_also_respected_and_boundary_is_allowed(self):
        class FakeTokenizer:
            model_max_length = 4

            def __call__(self, texts, **kwargs):
                return {"input_ids": [[1, 2, 3, 4] for _ in texts]}

        class FakeModel:
            max_seq_length = 128
            tokenizer = FakeTokenizer()

        report = stage1.guard_token_lengths(FakeModel(), [{"text": "Text", "variant_id": "A::basis"}])
        self.assertEqual(report["effective_token_limit"], 4)
        self.assertEqual(report["token_counts"], [4])


@unittest.skipUnless(importlib.util.find_spec("numpy") is not None, "NumPy fehlt, reine Datenprüfung bleibt verfügbar")
class SyntheticAggregationTests(unittest.TestCase):
    def setUp(self):
        import numpy as np
        self.np = np
        self.records = [
            {"item_id": "A", "array_group": "regular"},
            {"item_id": "A", "array_group": "regular"},
            {"item_id": "B", "array_group": "regular"},
            {"item_id": "K", "array_group": "controls"},
        ]

    def test_means_use_unit_formulations_then_item_normalization(self):
        np = self.np
        raw = np.array([[2., 0.], [0., 4.], [3., 4.]], dtype=np.float32)
        controls = np.array([[0., -7.]], dtype=np.float32)
        arrays = stage1.build_arrays(raw, controls, self.records, 2)
        np.testing.assert_array_equal(arrays["regular_raw"], raw)
        np.testing.assert_allclose(arrays["regular_l2"], [[1., 0.], [0., 1.], [.6, .8]])
        np.testing.assert_allclose(arrays["item_mean_of_raw"][0], [1., 2.])
        np.testing.assert_allclose(arrays["item_mean_of_l2"][0], [.5, .5])
        np.testing.assert_allclose(arrays["item_mean_of_l2_l2"][0], [2**-.5, 2**-.5])
        np.testing.assert_allclose(arrays["controls_l2"], [[0., -1.]])
        self.assertEqual(arrays["item_mean_of_l2"].shape, (2, 2))

    def test_zero_tiny_nonfinite_and_cancelling_means_fail(self):
        np = self.np
        for bad in ([[0., 0.]], [[1e-13, 0.]], [[float("nan"), 1.]], [[float("inf"), 1.]]):
            with self.subTest(bad=bad), self.assertRaises(stage1.ValidationError):
                stage1.l2_normalize(np.array(bad))
        with self.assertRaisesRegex(stage1.ValidationError, "L2-Norm"):
            stage1.build_arrays(np.array([[1., 0.], [-1., 0.], [1., 0.]]),
                                np.array([[1., 1.]]), self.records, 2)

    def test_wrong_dimension_and_control_count_are_rejected(self):
        np = self.np
        with self.assertRaisesRegex(stage1.ValidationError, "Falsche Embedding-Form"):
            stage1.build_arrays(np.ones((3, 3)), np.ones((1, 2)), self.records, 2)
        with self.assertRaisesRegex(stage1.ValidationError, "controls"):
            stage1.build_arrays(np.ones((3, 2)), np.ones((2, 2)), self.records, 2)

    def test_empty_control_array_is_supported_without_reclassifying_last_item(self):
        np = self.np
        arrays = stage1.build_arrays(np.ones((3, 2)), np.empty((0, 2)), self.records[:-1], 2)
        self.assertEqual(arrays["controls_raw"].shape, (0, 2))
        self.assertEqual(arrays["item_mean_of_l2_l2"].shape, (2, 2))


if __name__ == "__main__":
    unittest.main()
