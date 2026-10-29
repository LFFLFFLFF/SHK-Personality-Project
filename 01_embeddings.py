#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Stufe 1: Excel prüfen und nur auf ausdrücklichen Aufruf Embeddings berechnen.

Der Standardlauf benötigt ausschließlich openpyxl und die Standardbibliothek.
Textuelle Invertierung wird aus der Arbeitsfassung übernommen. Eine numerische
Invertierung von Vektoren oder das Ersetzen fehlender Daten findet niemals statt.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import csv
import hashlib
import importlib.metadata
import importlib.util
import json
from pathlib import Path
import platform
import re
import sys
import unicodedata
import uuid


SCRIPT_DIR = Path(__file__).resolve().parent
ITEM_HEADERS = (
    "Itemcode", "Originaltext", "Arbeitsfassung für die Analyse", "Fragebogen",
    "Zielkonstrukt", "Polung Original", "Variante 1", "Variante 2", "Variante 3",
    "Anmerkung", "Polung Analyse",
)
CONTROL_HEADERS = ("ID", "Rolle", "Text", "Anmerkung")
MAPPING_FIELDS = (
    "mapping_index", "array_group", "array_index", "item_index", "item_id",
    "variant_id", "role", "formulation_type", "construct", "questionnaire",
    "original_polarity", "analysis_polarity", "source_sheet", "source_table",
    "source_row", "source_column", "source_cell", "original_source_cell",
    "analysis_source_cell", "basis_source", "text", "text_sha256", "token_count",
)
LIBRARIES = {
    "openpyxl": "openpyxl", "numpy": "numpy", "torch": "torch",
    "sentence-transformers": "sentence_transformers", "transformers": "transformers",
    "tokenizers": "tokenizers", "huggingface-hub": "huggingface_hub",
    "safetensors": "safetensors",
}


class ValidationError(ValueError):
    """Eingaben werden bei unklarer Bedeutung nicht automatisch repariert."""


def sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _fingerprint(value):
    data = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(data.encode("utf-8")).hexdigest()


def _path(root, value):
    candidate = Path(value)
    return candidate.resolve() if candidate.is_absolute() else (Path(root) / candidate).resolve()


def load_config(path="config_embeddings.json"):
    """Relative Konfigurationspfade beziehen sich auf den Skriptordner, nie auf CWD."""
    config_path = _path(SCRIPT_DIR, path)
    try:
        with config_path.open(encoding="utf-8-sig") as stream:
            config = json.load(stream)
    except (OSError, ValueError) as exc:
        raise ValidationError(f"Konfiguration nicht lesbar: {config_path}: {exc}") from exc
    if not isinstance(config, dict) or config.get("schema_version") != 1:
        raise ValidationError("Die Konfiguration muss schema_version 1 besitzen.")
    required = (
        "input_file", "item_sheet", "item_table", "control_sheet", "control_table",
        "expected_items", "expected_variants", "expected_regular_texts",
        "expected_construct_counts", "expected_source_counts", "expected_control_count",
        "forbidden_item_ids", "expected_inverted_item_ids", "model_name", "model_revision",
        "expected_embedding_dimension", "embedding_prompt", "batch_size", "device", "output_dir",
    )
    missing = [key for key in required if key not in config]
    if missing:
        raise ValidationError("Fehlende Konfigurationsschlüssel: " + ", ".join(missing))
    for key in ("expected_items", "expected_variants", "expected_regular_texts",
                "expected_control_count", "expected_embedding_dimension", "batch_size"):
        if type(config[key]) is not int or config[key] < (1 if key in (
            "expected_items", "expected_embedding_dimension", "batch_size") else 0):
            raise ValidationError(f"{key} muss eine zulässige ganze Zahl sein.")
    for key in ("expected_construct_counts", "expected_source_counts"):
        if not isinstance(config[key], dict) or any(
            not isinstance(name, str) or type(count) is not int or count < 0
            for name, count in config[key].items()
        ):
            raise ValidationError(f"{key} muss Namen auf nichtnegative Anzahlen abbilden.")
    for key in ("forbidden_item_ids", "expected_inverted_item_ids"):
        values = config[key]
        if not isinstance(values, list) or any(not isinstance(v, str) or not v or
                v != v.strip() or re.search(r"[^\S ]| {2,}", v) for v in values) or len(set(values)) != len(values):
            raise ValidationError(f"{key} enthält ungültige oder doppelte IDs.")
    for key in ("input_file", "item_sheet", "item_table", "control_sheet", "control_table",
                "model_name", "model_revision", "device", "output_dir"):
        if not isinstance(config[key], str) or not config[key] or config[key] != config[key].strip():
            raise ValidationError(f"{key} muss eine nichtleere Zeichenkette ohne Rand-Leerraum sein.")
    if config["embedding_prompt"] != "":
        raise ValidationError("Stufe 1 verlangt embedding_prompt als explizite leere Zeichenkette.")
    if config["device"] != "cpu":
        raise ValidationError("Diese vorbereitete Stufe 1 ist ausdrücklich auf device cpu begrenzt.")
    if not re.fullmatch(r"[0-9a-fA-F]{40}", config["model_revision"]):
        raise ValidationError("model_revision muss ein vollständiger Commit-SHA mit 40 Hexzeichen sein.")
    return config, config_path


def library_status():
    """Paketmetadaten und Auffindbarkeit prüfen, ohne ML-Module zu importieren."""
    result = {}
    for distribution, module in LIBRARIES.items():
        try:
            version = importlib.metadata.version(distribution)
        except importlib.metadata.PackageNotFoundError:
            version = None
        try:
            available = importlib.util.find_spec(module) is not None
        except (ImportError, ValueError, AttributeError):
            available = False
        result[distribution] = {"version": version, "available": available}
    return result


def _cell_text(cell, *, optional=False, metadata=False, identifier=False):
    location = f"{cell.parent.title}!{cell.coordinate}"
    if cell.data_type in ("f", "e"):
        raise ValidationError(f"{location}: Formeln und Excel-Fehler sind nicht als Eingabe zulässig.")
    value = cell.value
    if value is None or value == "":
        if optional:
            return ""
        raise ValidationError(f"{location}: Pflichtfeld ist leer.")
    if not isinstance(value, str):
        raise ValidationError(f"{location}: Erwartet wird Text, erhalten wurde {type(value).__name__}.")
    text = value.strip()
    if not text:
        if optional:
            return ""
        raise ValidationError(f"{location}: Pflichtfeld enthält nur Leerraum.")
    if any((unicodedata.category(c) in ("Cf", "Cs") or
            (unicodedata.category(c) == "Cc" and c not in "\t\n\r") or c == "\ufffd") for c in text):
        raise ValidationError(f"{location}: Unsichtbare Steuerzeichen oder Ersatzzeichen im Text.")
    if text.casefold() in {"nan", "none", "null", "n/a", "#n/a", "-", "–", "—"}:
        raise ValidationError(f"{location}: Platzhalter ist kein gültiger Eingabetext.")
    if metadata and value != text:
        raise ValidationError(f"{location}: Metadaten dürfen keinen Rand-Leerraum enthalten.")
    # Original-IDs wie „RC3 36“ und „STR 29“ enthalten ein einzelnes inneres Leerzeichen.
    if identifier and re.search(r"[^\S ]| {2,}", text):
        raise ValidationError(f"{location}: IDs dürfen nur einfache innere Leerzeichen enthalten, keine übrigen Leerraumzeichen.")
    return text


def _table_rows(workbook, sheet_name, table_name, required_headers):
    if sheet_name not in workbook.sheetnames:
        raise ValidationError(f"Pflichtblatt fehlt: {sheet_name}.")
    sheet = workbook[sheet_name]
    if table_name not in sheet.tables:
        raise ValidationError(f"Native Excel-Tabelle fehlt: {sheet_name}/{table_name}.")
    from openpyxl.utils.cell import range_boundaries
    table = sheet.tables[table_name]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    names = [_cell_text(sheet.cell(min_row, col), metadata=True)
             for col in range(min_col, max_col + 1)]
    if len(set(names)) != len(names):
        raise ValidationError(f"{sheet_name}/{table_name}: Doppelte Spaltenüberschriften.")
    missing = sorted(set(required_headers) - set(names))
    if missing:
        raise ValidationError(f"{sheet_name}/{table_name}: Fehlende Spalten: {', '.join(missing)}")
    columns = dict(zip(names, range(min_col, max_col + 1)))
    # Ausschließlich table.ref ist maßgeblich. Fußnoten außerhalb werden nicht eingelesen.
    last_row = max_row - int(table.totalsRowCount or 0)
    rows = [{name: sheet.cell(row, col) for name, col in columns.items()}
            for row in range(min_row + 1, last_row + 1)]
    return rows, {"sheet": sheet_name, "table": table_name, "ref": table.ref,
                  "header_row": min_row, "data_rows": len(rows)}


def _comparison_key(text):
    # Nur der Duplikatvergleich ignoriert Großschreibung und mehrfachen Leerraum.
    # Die tatsächlich zu verarbeitenden Texte werden dabei nicht verändert.
    return " ".join(text.casefold().split())


def _record(cell, item_id, variant_id, role, formulation_type, **metadata):
    text = _cell_text(cell)
    return {"item_id": item_id, "variant_id": variant_id, "role": role,
            "formulation_type": formulation_type, "text": text,
            "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
            "source_sheet": cell.parent.title, "source_row": cell.row,
            "source_column": cell.column_letter, "source_cell": cell.coordinate,
            "token_count": "", **metadata}


def _read_workbook(config, input_path):
    import openpyxl
    workbook = openpyxl.load_workbook(input_path, read_only=False, data_only=False, keep_links=False)
    try:
        rows, item_source = _table_rows(workbook, config["item_sheet"], config["item_table"], ITEM_HEADERS)
        control_rows, control_source = _table_rows(
            workbook, config["control_sheet"], config["control_table"], CONTROL_HEADERS)
        records, fingerprints, seen_ids = [], [], set()
        inverse_ids, constructs, source_counts = set(), Counter(), Counter()
        for item_index, row in enumerate(rows):
            # Auch Anmerkungen und weitere Tabellenfelder dürfen keine Formeln/Fehler enthalten.
            values = {name: _cell_text(cell, optional=name in (
                "Arbeitsfassung für die Analyse", "Variante 3", "Anmerkung"))
                for name, cell in row.items()}
            item_id = _cell_text(row["Itemcode"], metadata=True, identifier=True)
            if item_id in seen_ids:
                raise ValidationError(f"Doppelte ID: {item_id}.")
            seen_ids.add(item_id)
            if item_id in config["forbidden_item_ids"]:
                raise ValidationError(f"Ausgeschlossenes Item ist enthalten: {item_id}.")
            construct = _cell_text(row["Zielkonstrukt"], metadata=True)
            questionnaire = _cell_text(row["Fragebogen"], metadata=True)
            original = _cell_text(row["Polung Original"], metadata=True)
            analysis = _cell_text(row["Polung Analyse"], metadata=True)
            if original not in ("direkt", "invertiert") or analysis != "direkt":
                raise ValidationError(f"{item_id}: Unzulässige Polung. Polung Analyse muss direkt sein.")
            is_inverse = original == "invertiert" or item_id.endswith("inv.")
            if is_inverse:
                inverse_ids.add(item_id)
                if item_id not in config["expected_inverted_item_ids"] or original != "invertiert":
                    raise ValidationError(f"{item_id}: Nicht erwartete oder widersprüchliche Invertierungsmetadaten.")
                if not values["Arbeitsfassung für die Analyse"]:
                    raise ValidationError(f"{item_id}: Invertiertes Item benötigt die textuell invertierte Arbeitsfassung.")
            basis_name = "Arbeitsfassung für die Analyse" if values["Arbeitsfassung für die Analyse"] else "Originaltext"
            basis = values[basis_name]
            source_counts["C" if basis_name == "Arbeitsfassung für die Analyse" else "B"] += 1
            variant_count = 3 if len(basis.split()) <= 25 else 2
            if (variant_count == 3 and not values["Variante 3"]) or (variant_count == 2 and values["Variante 3"]):
                raise ValidationError(f"{item_id}: Bei {len(basis.split())} Wörtern werden genau {variant_count} Varianten erwartet.")
            group_cells = [("basis", basis_name)] + [(f"v{n}", f"Variante {n}") for n in range(1, variant_count + 1)]
            group_texts = [_cell_text(row[name]) for _, name in group_cells]
            if len({_comparison_key(t) for t in group_texts}) != len(group_texts):
                raise ValidationError(f"{item_id}: Doppelte Formulierung innerhalb derselben Itemgruppe.")
            if is_inverse and any(_comparison_key(t) == _comparison_key(values["Originaltext"]) for t in group_texts):
                raise ValidationError(f"{item_id}: Gegenpoliger Originaltext darf nicht als Analyseformulierung eingehen.")
            common = {"array_group": "regular", "item_index": item_index,
                      "construct": construct, "questionnaire": questionnaire,
                      "original_polarity": original, "analysis_polarity": analysis,
                      "source_table": config["item_table"],
                      "original_source_cell": row["Originaltext"].coordinate,
                      "analysis_source_cell": row["Arbeitsfassung für die Analyse"].coordinate,
                      "basis_source": "C" if basis_name == "Arbeitsfassung für die Analyse" else "B"}
            for variant_id, name in group_cells:
                records.append(_record(row[name], item_id, f"{item_id}::{variant_id}", "item",
                                       "basis" if variant_id == "basis" else "variante", **common))
            constructs[construct] += 1
            fingerprints.append({"item_id": item_id, "source_row": row["Itemcode"].row,
                                 "sha256": _fingerprint(values)})
        if inverse_ids != set(config["expected_inverted_item_ids"]):
            raise ValidationError("Invertierte Item-IDs stimmen nicht exakt mit der erwarteten Menge überein.")
        for row in control_rows:
            for name, cell in row.items():
                _cell_text(cell, optional=name == "Anmerkung")
            item_id = _cell_text(row["ID"], metadata=True, identifier=True)
            role = _cell_text(row["Rolle"], metadata=True)
            if role not in ("kontrolle", "grenzfall"):
                raise ValidationError(f"{item_id}: Rolle muss explizit kontrolle oder grenzfall sein.")
            if item_id in seen_ids:
                raise ValidationError(f"Doppelte ID zwischen Item-/Kontrollrollen: {item_id}.")
            seen_ids.add(item_id)
            records.append(_record(row["Text"], item_id, f"{item_id}::text", role, "kontrolltext",
                array_group="controls", item_index="", construct="", questionnaire="",
                original_polarity="", analysis_polarity="", source_table=config["control_table"],
                original_source_cell="", analysis_source_cell="", basis_source=""))
        indices = Counter()
        for mapping_index, record in enumerate(records):
            group = record["array_group"]
            record.update(mapping_index=mapping_index, array_index=indices[group])
            indices[group] += 1
        return records, {"items": len(rows), "variants": sum(r["formulation_type"] == "variante" for r in records),
            "regular_texts": indices["regular"], "control_texts": indices["controls"], "total_texts": len(records),
            "construct_counts": dict(constructs), "source_counts": dict(source_counts),
            "control_role_counts": dict(Counter(r["role"] for r in records if r["array_group"] == "controls"))}, \
            fingerprints, [item_source, control_source]
    finally:
        workbook.close()


def validate_input(config, config_root):
    """Datenprüfung ohne Modell, NumPy, Torch, Netzwerk oder Embedding-Berechnung."""
    input_path = _path(config_root, config["input_file"])
    report = {"schema_version": 1, "stage": "stage1_embeddings", "mode": "pruefen",
        "created_at_utc": datetime.now(timezone.utc).isoformat(), "data_suitable": False,
        "model_loaded": False, "embeddings_computed": False, "model_downloads_allowed": False,
        "errors": [], "warnings": [], "counts": {}, "source_fingerprints": [],
        "input_file": str(input_path), "source_sha256": None, "source_tables": [],
        "script_sha256": sha256_file(__file__), "provenance": config.get("provenance", {}),
        "model": {"name": config["model_name"], "revision_sha": config["model_revision"],
                  "expected_dimension": config["expected_embedding_dimension"], "prompt": "",
                  "batch_size": config["batch_size"], "device": config["device"]},
        "software": {"python": sys.version, "platform": platform.platform(), "libraries": library_status()},
        "methodology": {"text_processing": "Nur äußeren Leerraum entfernen, innere Zeichen unverändert lassen.",
                        "basis_selection": "Nichtleere Arbeitsfassung, sonst Originaltext. Kein Datenersatz.",
                        "variant_rule": "Bis 25 Wörter drei Varianten, über 25 Wörter zwei Varianten.",
                        "inversion": "Textuelle Invertierung aus der Arbeitsfassung, niemals numerische Invertierung.",
                        "direction_check_limit": "Metadatenprüfung ersetzt keine inhaltliche Prüfung der Konstruktgleichrichtung.",
                        "duplicate_comparison": "Großschreibung und mehrfachen Leerraum nur im Vergleich ignorieren.",
                        "control_handling": "Explizite Rollen in separater Tabelle, keine automatische letzte Kontrollgruppe."},
        "stage_flags": {"squid_computed": False, "axes_computed": False, "cosine_computed": False}}
    for name, status in report["software"]["libraries"].items():
        if not status["available"] or status["version"] is None:
            report["warnings"].append(f"Paket {name} fehlt oder seine Version ist nicht ermittelbar. Datenprüfung bleibt möglich.")
    records = []
    try:
        if not input_path.is_file():
            raise ValidationError(f"Eingabedatei fehlt: {input_path}. Es gibt keine Ersatzdaten.")
        report["source_sha256"] = sha256_file(input_path)
        records, counts, fingerprints, sources = _read_workbook(config, input_path)
        report.update(counts=counts, source_fingerprints=fingerprints, source_tables=sources)
        expectations = {"items": config["expected_items"], "variants": config["expected_variants"],
                        "regular_texts": config["expected_regular_texts"],
                        "control_texts": config["expected_control_count"],
                        "construct_counts": config["expected_construct_counts"],
                        "source_counts": {k: v for k, v in config["expected_source_counts"].items() if v}}
        report["expected_counts"] = expectations
        for key, expected in expectations.items():
            if counts[key] != expected:
                report["errors"].append(f"Anzahlprüfung {key}: erwartet {expected}, gefunden {counts[key]}.")
        text_ids = defaultdict(set)
        for record in records:
            text_ids[_comparison_key(record["text"])].add(record["item_id"])
        shared = [sorted(ids) for ids in text_ids.values() if len(ids) > 1]
        if shared:
            report["warnings"].append(f"{len(shared)} Formulierungen werden unter unterschiedlichen IDs verwendet. Keine Deduplizierung.")
            report["cross_id_duplicate_texts"] = shared
    except (ValidationError, OSError, ValueError, KeyError, ImportError) as exc:
        report["errors"].append(str(exc))
    report["data_suitable"] = not report["errors"]
    return report, records


def write_report(directory, report, records):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    with (directory / "pruefbericht.json").open("w", encoding="utf-8") as stream:
        json.dump(report, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
    with (directory / "text_mapping.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=MAPPING_FIELDS, extrasaction="raise")
        writer.writeheader()
        writer.writerows(records)


def guard_token_lengths(model, records):
    """Jeden tatsächlichen Eingabetext vor dem ersten encode ohne Trunkierung tokenisieren."""
    limits = [int(model.max_seq_length)]
    tokenizer_limit = getattr(model.tokenizer, "model_max_length", None)
    if isinstance(tokenizer_limit, int) and 0 < tokenizer_limit < 10**9:
        limits.append(tokenizer_limit)
    limit = min(limits)
    if limit <= 0:
        raise ValidationError("Das Modell meldet keine gültige maximale Sequenzlänge.")
    tokenized = model.tokenizer([r["text"] for r in records], add_special_tokens=True,
        truncation=False, padding=False, return_attention_mask=False, return_token_type_ids=False)
    ids = tokenized["input_ids"]
    if len(ids) != len(records):
        raise ValidationError("Tokenizer-Ausgabe passt nicht zur Zahl der Eingabetexte.")
    lengths = [len(tokens) for tokens in ids]
    too_long = []
    for record, length in zip(records, lengths):
        record["token_count"] = length
        if length > limit:
            too_long.append(f"{record['variant_id']} ({length} > {limit})")
    if too_long:
        raise ValidationError("Tokenlimit überschritten. Keine stille Kürzung erlaubt: " + ", ".join(too_long))
    return {"model_max_seq_length": int(model.max_seq_length), "effective_token_limit": limit,
            "add_special_tokens": True, "truncation": False, "token_counts": lengths}


def l2_normalize(array):
    import numpy as np
    values = np.asarray(array)
    if values.ndim != 2 or values.dtype.kind != "f" or not np.isfinite(values).all():
        raise ValidationError("Nur endliche zweidimensionale Gleitkomma-Arrays sind zulässig.")
    norms = np.linalg.norm(values.astype(np.float64), axis=1, keepdims=True)
    if not np.isfinite(norms).all() or np.any(norms < 1e-12):
        raise ValidationError("L2-Norm ist null, kleiner als 1e-12 oder nicht endlich.")
    normalized = (values / norms).astype(values.dtype, copy=False)
    if not np.isfinite(normalized).all():
        raise ValidationError("Nichtendliche Werte nach der L2-Normierung.")
    return normalized


def build_arrays(regular_raw, controls_raw, records, dimension, *, save_raw_group_means=True):
    """Itemmittel aus gleich gewichteten, zunächst einzeln L2-normierten Formulierungen."""
    import numpy as np
    regular = [r for r in records if r["array_group"] == "regular"]
    controls = [r for r in records if r["array_group"] == "controls"]
    raw, control_raw = np.asarray(regular_raw), np.asarray(controls_raw)
    for name, array, count in (("regular", raw, len(regular)), ("controls", control_raw, len(controls))):
        if array.shape != (count, dimension):
            raise ValidationError(f"Falsche Embedding-Form für {name}: {array.shape}, erwartet {(count, dimension)}.")
    normed, control_normed = l2_normalize(raw), l2_normalize(control_raw)
    groups = defaultdict(list)
    for row, record in enumerate(regular):
        groups[record["item_id"]].append(row)
    mean_l2 = np.stack([normed[rows].mean(axis=0) for rows in groups.values()])
    arrays = {"regular_raw": raw, "regular_l2": normed,
              "item_mean_of_l2": mean_l2, "item_mean_of_l2_l2": l2_normalize(mean_l2),
              "controls_raw": control_raw, "controls_l2": control_normed}
    if save_raw_group_means:
        arrays["item_mean_of_raw"] = np.stack([raw[rows].mean(axis=0) for rows in groups.values()])
    if any(not np.isfinite(array).all() for array in arrays.values()):
        raise ValidationError("Nichtendliche Werte in einem zu speichernden Array.")
    return arrays


def compute_embeddings(config, config_root, report, records):
    """Einziger Einstieg für Modellladen und Inferenz. Aufruf nur über --berechnen."""
    if not report["data_suitable"]:
        raise ValidationError("Ungültige Daten dürfen nicht eingebettet werden.")
    if not re.fullmatch(r"[0-9a-fA-F]{40}", config["model_revision"]):
        raise ValidationError("Ohne festen Commit-SHA wird kein Modell geladen.")
    report.update(mode="berechnen", model_downloads_allowed=True)
    import numpy as np
    import torch
    from sentence_transformers import SentenceTransformer
    torch.manual_seed(0)
    model = SentenceTransformer(config["model_name"], revision=config["model_revision"],
        trust_remote_code=False, device=config["device"], prompts={}, default_prompt_name=None)
    model.prompts = {}
    model.default_prompt_name = None
    model.eval()
    report["model_loaded"] = True
    report["model"].update(prompts={}, default_prompt_name=None, trust_remote_code=False)
    first_module = model._first_module()
    resolved_sha = getattr(getattr(getattr(first_module, "auto_model", None), "config", None), "_commit_hash", None)
    report["model"]["resolved_revision_sha"] = resolved_sha
    if resolved_sha and resolved_sha.lower() != config["model_revision"].lower():
        raise ValidationError("Geladener Modell-Commit stimmt nicht mit dem festgelegten Commit überein.")
    report["tokenization"] = guard_token_lengths(model, records)
    inference_calls, raw_arrays = [], {}
    for group in ("regular", "controls"):
        texts = [r["text"] for r in records if r["array_group"] == group]
        if texts:
            raw_arrays[group] = model.encode(texts, batch_size=config["batch_size"], prompt="",
                normalize_embeddings=False, convert_to_numpy=True, convert_to_tensor=False,
                show_progress_bar=True)
            inference_calls.append({"array_group": group, "text_count": len(texts),
                "batch_size": config["batch_size"], "prompt": "", "normalize_embeddings": False})
        else:
            raw_arrays[group] = np.empty((0, config["expected_embedding_dimension"]), dtype=np.float32)
    arrays = build_arrays(raw_arrays["regular"], raw_arrays["controls"], records,
        config["expected_embedding_dimension"], save_raw_group_means=config.get("save_raw_group_means", True))
    report.update(embeddings_computed=True, inference_calls=inference_calls,
                  inference_call_count=len(inference_calls))
    output_root = _path(config_root, config["output_dir"])
    output_root.mkdir(parents=True, exist_ok=True)
    run_name = datetime.now(timezone.utc).strftime("run_%Y%m%dT%H%M%S_%fZ_") + report["source_sha256"][:8] + "_" + uuid.uuid4().hex[:8]
    run_dir = output_root / run_name
    run_dir.mkdir(exist_ok=False)
    report["run_directory"] = str(run_dir)
    report["array_outputs"] = {}
    for name, array in arrays.items():
        destination = run_dir / f"{name}.npy"
        np.save(destination, array, allow_pickle=False)
        report["array_outputs"][name] = {"file": destination.name, "shape": list(array.shape),
            "dtype": str(array.dtype), "sha256": sha256_file(destination)}
    report["installed_distributions"] = dict(sorted(
        (distribution.metadata.get("Name", "unbekannt"), distribution.version)
        for distribution in importlib.metadata.distributions()))
    report["item_row_mapping"] = [{"array_index": r["item_index"], "item_id": r["item_id"]}
                                   for r in records if r["formulation_type"] == "basis"]
    write_report(run_dir, report, records)
    report["mapping_sha256"] = sha256_file(run_dir / "text_mapping.csv")
    with (run_dir / "manifest.json").open("w", encoding="utf-8") as stream:
        json.dump(report, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
    return run_dir


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", default="config_embeddings.json", help="Relativ zum Skriptordner oder absolut")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--pruefen", action="store_true", help="Ausschließlich prüfen. Auch ohne Modus ist dies der Standard.")
    mode.add_argument("--berechnen", action="store_true", help="Nach erfolgreicher Prüfung ausdrücklich Modell laden und Embeddings berechnen")
    parser.add_argument("--bericht-dir", help="Optionaler Prüfberichtordner, relativ zur Konfiguration oder absolut")
    args = parser.parse_args(argv)
    try:
        config, config_path = load_config(args.config)
    except ValidationError as exc:
        print(f"FEHLER: {exc}", file=sys.stderr)
        return 2
    report, records = validate_input(config, config_path.parent)
    report.update(config_file=str(config_path), config_sha256=sha256_file(config_path))
    try:
        if args.berechnen and report["data_suitable"]:
            run_dir = compute_embeddings(config, config_path.parent, report, records)
            print(f"Stufe 1 gespeichert: {run_dir}")
    except Exception as exc:
        report["computation_error"] = f"{type(exc).__name__}: {exc}"
        print(f"FEHLER bei ausdrücklich angeforderter Berechnung: {report['computation_error']}", file=sys.stderr)
    finally:
        if args.bericht_dir:
            write_report(_path(config_path.parent, args.bericht_dir), report, records)
    for error in report["errors"]:
        print(f"FEHLER: {error}", file=sys.stderr)
    for warning in report["warnings"]:
        print(f"WARNUNG: {warning}")
    if not report["data_suitable"] or "computation_error" in report:
        return 2
    counts = report["counts"]
    print(f"Datenprüfung erfolgreich: {counts['items']} Items, {counts['variants']} Varianten, "
          f"{counts['regular_texts']} reguläre Texte, {counts['control_texts']} separate Kontroll-/Grenzfalltexte.")
    if not args.berechnen:
        print("Kein Modell geladen. Keine Embeddings berechnet. Keine Downloads gestartet.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
